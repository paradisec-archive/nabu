#!/usr/bin/env -S uv run --quiet --with openpyxl python3
"""
PROTOTYPE (#1185) — candidate ISO 639-3 / Glottolog / AUSTLANG concordance for Nabu's languages.

Throwaway. Joins Nabu's languages to Glottolog by ISO code and to AUSTLANG by three independent
routes (Bowern's Chirila code table, describo data-packs, normalised name/synonym matching), and
writes one spreadsheet for Nick to work through with every ambiguity left visible.

Run from the repo root:

    bin/nabu_run bin/rails runner docs/research/language-concordance/export_nabu_languages.rb
    cp tmp/nabu-languages.csv tmp/language-concordance/
    cp iso-639-3_Code_Tables_*.zip tmp/language-concordance/      # from https://iso639-3.sil.org/code_tables/download_tables
    docs/research/language-concordance/concordance.py

Every other source is downloaded into tmp/language-concordance/ on first run.
"""
import csv
import difflib
import io
import json
import math
import re
import sys
import unicodedata
import urllib.request
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]
DATA = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'tmp' / 'language-concordance'
OUT = Path(__file__).resolve().parent
DATA.mkdir(parents=True, exist_ok=True)

SOURCES = {
    'glottolog-languages.csv': 'https://raw.githubusercontent.com/glottolog/glottolog-cldf/master/cldf/languages.csv',
    'austlang.json': 'https://data.gov.au/data/api/3/action/datastore_search?resource_id=e9a9ea06-d821-4b53-a05f-877409a1a19c&limit=2000',
    'describo-austlang.json': 'https://raw.githubusercontent.com/describo/data-packs/master/data-packs/languages/austlang-language-data-pack.json',
    'describo-glottolog.json': 'https://raw.githubusercontent.com/describo/data-packs/master/data-packs/languages/glottolog-language-data-pack.json',
    'chirila-codes.xlsx': 'https://zenodo.org/api/records/4898185/files/Chirila%20Language%20Codes.xlsx/content',
}
SPECIAL = {'mul': 'multiple languages', 'und': 'undetermined', 'zxx': 'no linguistic content'}

def fetch(name):
    path = DATA / name
    if not path.exists():
        print(f'downloading {name}', file=sys.stderr)
        req = urllib.request.Request(SOURCES[name], headers={'User-Agent': 'nabu-concordance-prototype'})
        path.write_bytes(urllib.request.urlopen(req, timeout=120).read())
    return path

def norm(s):
    s = unicodedata.normalize('NFKD', s or '')
    s = ''.join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r'\(.*?\)|\[.*?\]', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s.lower())

def name_variants(s):
    out = {norm(s)}
    if ',' in s:
        parts = [p.strip() for p in s.split(',')]
        out.add(norm(' '.join(reversed(parts))))
    return {v for v in out if v}

def km(a, b):
    if not a or not b:
        return None
    (lat1, lon1), (lat2, lon2) = a, b
    p = math.pi / 180
    h = 0.5 - math.cos((lat2 - lat1) * p) / 2 + math.cos(lat1 * p) * math.cos(lat2 * p) * (1 - math.cos((lon2 - lon1) * p)) / 2
    return round(12742 * math.asin(math.sqrt(h)))

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None

# ---------------------------------------------------------------- load sources
nabu_csv = DATA / 'nabu-languages.csv'
if not nabu_csv.exists():
    sys.exit(f'{nabu_csv} missing: run export_nabu_languages.rb first (see header)')
nabu = list(csv.DictReader(open(nabu_csv, encoding='utf-8')))
for r in nabu:
    r['usage'] = int(r['usage_count'])
    r['in_use'] = r['usage'] > 0
    r['in_au'] = r['in_au'] == 'true'
    r['retired'] = r['retired'] in ('true', 't', '1')
    r['has_box'] = r['has_box'] == 'true'
    n, s, e, w = (num(r[k]) for k in ('north', 'south', 'east', 'west'))
    r['box_centre'] = ((n + s) / 2, (e + w) / 2) if r['has_box'] else None

glotto = list(csv.DictReader(open(fetch('glottolog-languages.csv'), encoding='utf-8')))
glotto_by_id = {g['ID']: g for g in glotto}
glotto_by_iso = {g['ISO639P3code']: g for g in glotto if g['ISO639P3code']}
glotto_names = defaultdict(set)
for g in glotto:
    glotto_names[norm(g['Name'])].add(g['ID'])
    g['point'] = (num(g['Latitude']), num(g['Longitude'])) if g['Latitude'] else None

austlang = {}
austlang_names = defaultdict(lambda: defaultdict(set))  # norm -> code -> {'name'|'synonym'}
austlang_norm_to_label = {}
for rec in json.load(open(fetch('austlang.json')))['result']['records']:
    code = rec['language_code'].strip()
    lat, lon = num(rec.get('approximate_latitude_of_language_variety')), num(rec.get('approximate_longitude_of_language_variety'))
    austlang[code] = {
        'code': code, 'name': rec['language_name'].strip(), 'synonyms': [s.strip() for s in (rec.get('language_synonym') or '').split('|') if s.strip()],
        'uri': rec.get('uri'), 'point': (lat, lon) if lat and lon else None,
    }
    for part in austlang[code]['name'].split('/'):
        austlang_names[norm(part)][code].add('name')
        austlang_norm_to_label.setdefault(norm(part), set()).add(f'{code} {part.strip()}')
    for syn in austlang[code]['synonyms']:
        austlang_names[norm(syn)][code].add('synonym')
        austlang_norm_to_label.setdefault(norm(syn), set()).add(f'{code} {austlang[code]["name"]} (synonym {syn})')

def austlang_code(raw):
    """Normalise a Chirila/describo AIATSIS code onto a code that exists in the datastore; keeps Bowern's trailing * as a flag."""
    c = str(raw).strip()
    star = '*' if c.endswith('*') else ''
    c = c.rstrip('*').strip()
    if c in austlang:
        return c + star
    base = re.sub(r'\..*$', '', c)
    return base + star if base in austlang else None

# ISO 639-3 tables
iso_zip = sorted(DATA.glob('iso-639-3_Code_Tables_*.zip')) + sorted(ROOT.glob('iso-639-3_Code_Tables_*.zip'))
if not iso_zip:
    sys.exit('no iso-639-3_Code_Tables_*.zip in data dir; download from https://iso639-3.sil.org/code_tables/download_tables')
iso_release = re.search(r'(\d{8})', iso_zip[-1].name).group(1)
with zipfile.ZipFile(iso_zip[-1]) as z:
    def tab(suffix):
        name = next(n for n in z.namelist() if n.endswith(suffix))
        return list(csv.DictReader(io.TextIOWrapper(z.open(name), encoding='utf-8'), delimiter='\t'))
    iso_info = {r['Id']: r for r in tab('iso-639-3.tab')}
    macro_members = defaultdict(list)
    for r in tab('iso-639-3-macrolanguages.tab'):
        macro_members[r['M_Id']].append((r['I_Id'], r['I_Status']))
    iso_retired = {r['Id']: r for r in tab('iso-639-3_Retirements.tab')}

# describo data-packs (their concordance is exact primary-name match, first hit on collision)
describo_glotto_to_austlang = defaultdict(set)
describo_austlang_to_glotto = defaultdict(set)
for e in json.load(open(fetch('describo-glottolog.json'))):
    if e.get('austlangCode'):
        c = austlang_code(e['austlangCode'])
        if c:
            describo_glotto_to_austlang[e['languageCode']].add(c)
for e in json.load(open(fetch('describo-austlang.json'))):
    if e.get('glottologCode'):
        c = austlang_code(e['languageCode'])
        if c:
            describo_austlang_to_glotto[c].add(e['glottologCode'])
            describo_glotto_to_austlang[e['glottologCode']].add(c)

# Chirila (Bowern 2021, Zenodo 4898185, CC-BY-NC-4.0): hand-curated glottocode / ISO / AIATSIS per language
chirila_by_iso = defaultdict(list)
chirila_by_glotto = defaultdict(list)
chirila_by_austlang = defaultdict(list)
ws = openpyxl.load_workbook(fetch('chirila-codes.xlsx'), read_only=True).worksheets[0]
rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
head = rows[0]
for r in rows[1:]:
    r = dict(zip(head, list(r) + [None] * (len(head) - len(r))))
    iso_raw = str(r.get('ISO639') or '').strip()
    iso_m = re.match(r'^(\?)?([a-z]{3})', iso_raw)
    entry = {
        'name': r.get('StandardLanguageName'), 'glottocode': (r.get('glottolog@id') or '').strip() or None,
        'iso': iso_m.group(2) if iso_m else None, 'iso_uncertain': bool(iso_m and iso_m.group(1)),
        'austlang': [(c.rstrip('*'), c.endswith('*')) for c in (austlang_code(x) for x in str(r.get('AIATSIS_Code') or '').split(',') if x.strip()) if c],
    }
    if entry['iso']:
        chirila_by_iso[entry['iso']].append(entry)
    if entry['glottocode']:
        chirila_by_glotto[entry['glottocode']].append(entry)
    for c, _ in entry['austlang']:
        chirila_by_austlang[c].append(entry)

# ---------------------------------------------------------------- per Nabu language
nabu_by_code = {r['code']: r for r in nabu}
nabu_names = defaultdict(set)
for r in nabu:
    for v in name_variants(r['name']):
        nabu_names[v].add(r['code'])

for r in nabu:
    code = r['code']
    g = glotto_by_iso.get(code)
    r['glottocode'] = g['ID'] if g else None
    r['glottolog_name'] = g['Name'] if g else None
    r['glottolog_level'] = g['Level'] if g else None
    r['glottolog_parent'] = g['Language_ID'] if g else None
    r['glottolog_closest_iso'] = g['Closest_ISO369P3code'] if g else None
    r['glottolog_point'] = g['point'] if g else None
    r['ref_point'] = r['glottolog_point'] or r['box_centre']
    r['australian'] = g['Macroarea'] == 'Australia' if g else r['in_au']

    # why no glottocode
    r['glottolog_gap'] = None
    r['glottolog_gap_detail'] = None
    if not g:
        if code in SPECIAL:
            r['glottolog_gap'] = 'special code'
            r['glottolog_gap_detail'] = f'ISO special code ({SPECIAL[code]}); no Glottolog equivalent by design'
        elif code in macro_members:
            r['glottolog_gap'] = 'macrolanguage'
            members = []
            for m, status in macro_members[code]:
                mg = glotto_by_iso.get(m)
                members.append(f"{m}={mg['ID'] if mg else '?'}" + ('' if status == 'A' else ' (retired)'))
            fam = ', '.join(sorted(glotto_names.get(norm(iso_info[code]['Ref_Name']), [])))
            r['glottolog_gap_detail'] = f"{len(members)} member languages: {' '.join(members)}" + (f'; Glottolog languoid named the same: {fam}' if fam else '')
        elif code in iso_retired:
            ret = iso_retired[code]
            to = ret['Change_To']
            tg = glotto_by_iso.get(to)
            r['glottolog_gap'] = 'retired in ISO'
            r['glottolog_gap_detail'] = f"retired {ret['Effective']} reason {ret['Ret_Reason']}" + (f"; changed to {to} ({tg['ID'] if tg else 'no glottocode'})" if to else f"; remedy: {' '.join(ret['Ret_Remedy'].split())}")
        elif code not in iso_info:
            r['glottolog_gap'] = 'not in ISO 639-3'
            r['glottolog_gap_detail'] = f'code absent from the {iso_release} ISO tables'
        else:
            cands = sorted(glotto_names.get(norm(r['name']), set()) | glotto_names.get(norm(iso_info[code]['Ref_Name']), set()))
            r['glottolog_gap'] = 'no glottocode for ISO'
            r['glottolog_gap_detail'] = f"ISO {iso_info[code]['Ref_Name']} (scope {iso_info[code]['Scope']}, type {iso_info[code]['Language_Type']}); " + \
                (f"Glottolog languoids with the same name: {', '.join(f'{c} ({glotto_by_id[c]['Level']})' for c in cands)}" if cands else 'no Glottolog languoid with the same name')

    # AUSTLANG candidates: code -> set(methods)
    cands = defaultdict(set)
    for entry in chirila_by_iso.get(code, []):
        for c, star in entry['austlang']:
            cands[c].add('chirila:iso' + ('?' if entry['iso_uncertain'] else '') + ('*' if star else ''))
    if r['glottocode']:
        for entry in chirila_by_glotto.get(r['glottocode'], []):
            for c, star in entry['austlang']:
                cands[c].add('chirila:glottocode' + ('*' if star else ''))
        for c in describo_glotto_to_austlang.get(r['glottocode'], []):
            cands[c.rstrip('*')].add('describo')
    if r['australian'] or cands:
        for v in name_variants(r['name']):
            for c, kinds in austlang_names.get(v, {}).items():
                cands[c].update('nabu-name' if k == 'name' else 'nabu-name→synonym' for k in kinds)
        if r['glottolog_name']:
            for c, kinds in austlang_names.get(norm(r['glottolog_name']), {}).items():
                cands[c].update('glottolog-name' if k == 'name' else 'glottolog-name→synonym' for k in kinds)
    r['australian'] = r['australian'] or bool(cands)
    r['austlang_candidates'] = {c: sorted(m) for c, m in cands.items()}
    for c in cands:
        austlang[c].setdefault('nabu_hits', {})[code] = sorted(cands[c])
    if not r['australian']:
        r['austlang_status'] = 'n/a'
    elif not cands:
        r['austlang_status'] = 'none'
    elif len(cands) == 1:
        r['austlang_status'] = 'single'
    else:
        r['austlang_status'] = 'ambiguous'
    curated = [c for c, m in cands.items() if any(x.startswith('chirila') for x in m)]
    r['austlang_suggested'] = curated[0] if len(curated) == 1 else None
    r['austlang_suggested_starred'] = bool(curated) and any(x.startswith('chirila') and x.endswith('*') for x in cands[curated[0]])

def cand_text(r):
    parts = []
    for c, methods in sorted(r['austlang_candidates'].items(), key=lambda kv: (-len(kv[1]), kv[0])):
        d = km(r['ref_point'], austlang[c]['point'])
        parts.append(f"{c} {austlang[c]['name']} [{', '.join(methods)}]" + (f' {d} km' if d is not None else ''))
    return '; '.join(parts)

# ---------------------------------------------------------------- write
BOLD = Font(bold=True)
FILL = PatternFill('solid', fgColor='FFF2CC')
wb = openpyxl.Workbook()
wb.remove(wb.active)

def sheet(title, header, rows, widths=None, note_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = BOLD
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h, min(60, max(10, len(h) + 2)))
        ws.column_dimensions[get_column_letter(i)].width = w
        if h in note_cols:
            for c in ws[get_column_letter(i)][1:]:
                c.fill = FILL
    return ws

def as_csv(name, header, rows):
    with open(OUT / name, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)

nabu.sort(key=lambda r: (not r['in_use'], not r['australian'], -r['usage'], r['code']))
in_use = [r for r in nabu if r['in_use']]
aus_in_use = [r for r in in_use if r['australian']]
gaps = [r for r in nabu if not r['glottocode']]
ambiguous = [r for r in aus_in_use if r['austlang_status'] == 'ambiguous']
unmatched = [r for r in aus_in_use if r['austlang_status'] == 'none']
single = [r for r in aus_in_use if r['austlang_status'] == 'single']
counts = {
    'Nabu languages': len(nabu), 'in use': len(in_use), 'retired': sum(r['retired'] for r in nabu),
    'with a glottocode': sum(1 for r in nabu if r['glottocode']), 'no glottocode': len(gaps), 'no glottocode, in use': sum(1 for r in gaps if r['in_use']),
    'Australian (Glottolog macroarea Australia; Nabu AU country link when Glottolog lacks the code)': sum(r['australian'] for r in nabu), 'Australian in use': len(aus_in_use),
    'Nabu AU country link but not Australian by the above (English, Italian, ...)': sum(1 for r in nabu if r['in_au'] and not r['australian']),
    'Australian in use: single AUSTLANG candidate': len(single),
    'Australian in use: single candidate backed by Chirila': sum(1 for r in single if r['austlang_suggested']),
    'Australian in use: ambiguous (2+ candidates)': len(ambiguous),
    'Australian in use: ambiguous but Chirila picks exactly one': sum(1 for r in ambiguous if r['austlang_suggested']),
    'Australian in use: no AUSTLANG candidate': len(unmatched),
    'AUSTLANG varieties': len(austlang),
    'AUSTLANG varieties reachable from some Nabu language': sum(1 for a in austlang.values() if a.get('nabu_hits')),
    'AUSTLANG varieties reachable from an in-use Nabu language': sum(1 for a in austlang.values() if any(nabu_by_code[c]['in_use'] for c in a.get('nabu_hits', {}))),
    'AUSTLANG varieties with no Nabu language at all': sum(1 for a in austlang.values() if not a.get('nabu_hits')),
}

readme = wb.create_sheet('README')
readme.column_dimensions['A'].width = 120
readme_lines = [
    'PROTOTYPE — candidate ISO 639-3 / Glottolog / AUSTLANG concordance for Nabu (issue #1185, map #1182)',
    f'Generated {date.today()} by docs/research/language-concordance/concordance.py. Nothing here is applied to Nabu; it is evidence for the questionnaire in #1186.',
    '',
    'HOW TO USE',
    '- "Nabu languages" is the main sheet: one row per Nabu language, sorted in-use Australian first, then other in-use, then the rest.',
    '- Yellow columns (decision, notes) are for Nick. Everything else is generated.',
    '- "No Glottolog match" lists every ISO code Glottolog has no languoid for, with why (special, macrolanguage, retired, orphan) and what it could map to.',
    '- "AUSTLANG ambiguous" is one row per candidate for each in-use Australian language with 2+ AUSTLANG candidates: tick one, or none.',
    '- "AUSTLANG unmatched" is each in-use Australian language with no candidate, with the nearest AUSTLANG varieties by distance as hints.',
    '- "AUSTLANG varieties" is the reverse view: all 1,204 varieties, which Nabu languages reach them, and which reach nothing (these would be new ISO-less rows).',
    '',
    'HOW CANDIDATES WERE FOUND (method tags in square brackets)',
    '- chirila:iso — Bowern 2021 "Chirila Language Codes" lists this AIATSIS code on the row with this ISO code (chirila:iso? = Bowern marked the ISO uncertain).',
    '- chirila:glottocode — same table, matched through the glottocode instead.',
    '- describo — LDaCA describo data-packs cross-reference (exact primary-name match, first hit on collision, so treat as weak).',
    '- nabu-name / glottolog-name — Nabu\'s name (or Glottolog\'s) equals an AUSTLANG primary name after normalising case, punctuation, diacritics and bracketed text.',
    '- …→synonym — same, but matched an AUSTLANG synonym rather than the primary name.',
    '- Distances are from the Glottolog point (or Nabu\'s bounding-box centre if Glottolog has none) to the AUSTLANG point. 369 AUSTLANG varieties have no usable point.',
    '- austlang_suggested is filled only when Chirila names exactly one candidate; it is a hint, not a decision. A trailing * (also on method tags) means Bowern starred that code in her table; the deposit does not say why.',
    '- australian = Glottolog macroarea Australia (Nabu AU country link only for codes Glottolog lacks), or a Chirila/describo link exists. Nabu\'s AU country link alone is not used because Ethnologue lists English, Italian and Mandarin under Australia; see nabu_au_country_link.',
    '',
    'COUNTS',
] + [f'- {k}: {v:,}' for k, v in counts.items()] + [
    '',
    'CAVEATS',
    '- AUSTLANG publishes no ISO code and no glottocode; every link is inferred. Chirila is the only human-curated source and covers ~370 languages.',
    '- Chirila is CC-BY-NC-4.0 (Bowern 2021, doi:10.5281/zenodo.4898185); fine as evidence for a decision, licence check before shipping data derived from it.',
    '- AUSTLANG is variety-level (1,204 entries) against ~450 Nabu ISO codes linked to Australia; many-to-many is expected, not an error.',
    '- Names in Nabu come from Ethnologue and have never been curated against AUSTLANG spellings.',
]
for i, line in enumerate(readme_lines, 1):
    readme.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True)
    if line.isupper() or i == 1:
        readme.cell(row=i, column=1).font = BOLD

header = ['nabu_id', 'iso639_3', 'nabu_name', 'retired', 'usage', 'in_use', 'australian', 'nabu_au_country_link', 'has_box',
          'glottocode', 'glottolog_name', 'glottolog_level', 'glottolog_closest_iso', 'glottolog_gap', 'glottolog_gap_detail',
          'austlang_status', 'austlang_suggested', 'austlang_suggested_name', 'austlang_candidates', 'decision', 'notes']
main_rows = [[r['id'], r['code'], r['name'], r['retired'], r['usage'], r['in_use'], r['australian'], r['in_au'], r['has_box'],
              r['glottocode'], r['glottolog_name'], r['glottolog_level'], r['glottolog_closest_iso'], r['glottolog_gap'], r['glottolog_gap_detail'],
              r['austlang_status'], (r['austlang_suggested'] + ('*' if r['austlang_suggested_starred'] else '')) if r['austlang_suggested'] else None, austlang[r['austlang_suggested']]['name'] if r['austlang_suggested'] else None,
              cand_text(r) or None, None, None] for r in nabu]
sheet('Nabu languages', header, main_rows, {'nabu_name': 30, 'glottolog_name': 24, 'glottolog_gap_detail': 60, 'austlang_candidates': 90, 'decision': 16, 'notes': 30}, ('decision', 'notes'))
as_csv('nabu-languages-concordance.csv', header, main_rows)

header = ['iso639_3', 'nabu_name', 'usage', 'retired_in_nabu', 'gap', 'detail', 'decision', 'notes']
gap_rows = [[r['code'], r['name'], r['usage'], r['retired'], r['glottolog_gap'], r['glottolog_gap_detail'], None, None]
            for r in sorted(gaps, key=lambda r: (r['glottolog_gap'], -r['usage'], r['code']))]
sheet('No Glottolog match', header, gap_rows, {'nabu_name': 30, 'detail': 110, 'decision': 16, 'notes': 30}, ('decision', 'notes'))
as_csv('no-glottolog-match.csv', header, gap_rows)

header = ['iso639_3', 'nabu_name', 'usage', 'glottocode', 'glottolog_name', 'n_candidates', 'austlang_code', 'austlang_name', 'austlang_synonyms', 'methods', 'km_from_reference', 'chirila_picks', 'choose', 'notes']
amb_rows = []
for r in ambiguous:
    for c, methods in sorted(r['austlang_candidates'].items(), key=lambda kv: (-len(kv[1]), kv[0])):
        a = austlang[c]
        amb_rows.append([r['code'], r['name'], r['usage'], r['glottocode'], r['glottolog_name'], len(r['austlang_candidates']), c, a['name'], ' | '.join(a['synonyms'])[:200],
                         ', '.join(methods), km(r['ref_point'], a['point']), r['austlang_suggested'] == c or None, None, None])
sheet('AUSTLANG ambiguous', header, amb_rows, {'nabu_name': 26, 'glottolog_name': 22, 'austlang_name': 30, 'austlang_synonyms': 50, 'methods': 40, 'choose': 10, 'notes': 30}, ('choose', 'notes'))
as_csv('austlang-ambiguous.csv', header, amb_rows)

header = ['iso639_3', 'nabu_name', 'usage', 'glottocode', 'glottolog_name', 'glottolog_level', 'reference_point', 'nearest_austlang_by_distance', 'similar_austlang_names', 'austlang_code', 'notes']
unm_rows = []
for r in unmatched:
    near = ''
    if r['ref_point']:
        ds = sorted((km(r['ref_point'], a['point']), a['code'], a['name']) for a in austlang.values() if a['point'])
        near = '; '.join(f'{c} {n} {d} km' for d, c, n in ds[:5])
    similar = []
    for q in [r['name'], r['glottolog_name'] or '']:
        for m in difflib.get_close_matches(norm(q), austlang_norm_to_label.keys(), n=4, cutoff=0.8):
            similar.extend(sorted(austlang_norm_to_label[m]))
    unm_rows.append([r['code'], r['name'], r['usage'], r['glottocode'], r['glottolog_name'], r['glottolog_level'],
                     f"{r['ref_point'][0]:.2f}, {r['ref_point'][1]:.2f}" if r['ref_point'] else None, near or None, '; '.join(dict.fromkeys(similar)) or None, None, None])
sheet('AUSTLANG unmatched', header, unm_rows, {'nabu_name': 26, 'glottolog_name': 22, 'nearest_austlang_by_distance': 90, 'similar_austlang_names': 70, 'austlang_code': 14, 'notes': 30}, ('austlang_code', 'notes'))
as_csv('austlang-unmatched.csv', header, unm_rows)

header = ['austlang_code', 'austlang_name', 'has_point', 'nabu_languages_reaching_it', 'nabu_usage_total', 'chirila_iso', 'chirila_glottocode', 'describo_glottocode', 'status', 'uri']
var_rows = []
for a in sorted(austlang.values(), key=lambda a: (re.sub(r'\d+', '', a['code']), int(re.sub(r'\D', '', a['code']) or 0))):
    hits = a.get('nabu_hits', {})
    usage = sum(nabu_by_code[c]['usage'] for c in hits)
    ch = chirila_by_austlang.get(a['code'], [])
    if not hits:
        status = 'no Nabu language (would be a new ISO-less row)'
    elif len(hits) == 1:
        status = 'one Nabu language'
    else:
        status = f'{len(hits)} Nabu languages'
    var_rows.append([a['code'], a['name'], bool(a['point']),
                     '; '.join(f"{c} {nabu_by_code[c]['name']} [{', '.join(m)}]" for c, m in sorted(hits.items())) or None, usage,
                     ', '.join(sorted({('?' if e['iso_uncertain'] else '') + e['iso'] for e in ch if e['iso']})) or None,
                     ', '.join(sorted({e['glottocode'] for e in ch if e['glottocode']})) or None,
                     ', '.join(sorted(describo_austlang_to_glotto.get(a['code'], []))) or None, status, a['uri']])
sheet('AUSTLANG varieties', header, var_rows, {'austlang_name': 34, 'nabu_languages_reaching_it': 90, 'status': 44, 'uri': 55})
as_csv('austlang-varieties.csv', header, var_rows)

sources = wb.create_sheet('Sources')
sources.column_dimensions['A'].width = 28
sources.column_dimensions['B'].width = 120
for row in [
    ('Nabu', f'tmp/nabu-languages.csv exported by export_nabu_languages.rb from the development database on {date.today()}; usage = collection_languages + item_content_languages + item_subject_languages'),
    ('Glottolog', f"{SOURCES['glottolog-languages.csv']} (CLDF languages.csv, master; {len(glotto):,} languoids). CC-BY-4.0"),
    ('AUSTLANG', f"{SOURCES['austlang.json']} ({len(austlang):,} records)"),
    ('ISO 639-3', f'SIL code tables release {iso_release}: iso-639-3.tab, iso-639-3-macrolanguages.tab, iso-639-3_Retirements.tab'),
    ('describo data-packs', 'github.com/describo/data-packs master: austlang-language-data-pack.json and glottolog-language-data-pack.json (cross-referenced by exact primary name, first hit)'),
    ('Chirila', f"{SOURCES['chirila-codes.xlsx']} — Bowern 2021 Files for Australian Language Locations, CC-BY-NC-4.0; {sum(len(v) for v in chirila_by_iso.values())} rows with an ISO code, {len(chirila_by_austlang)} distinct AIATSIS codes"),
]:
    sources.append(row)
for c in sources['A']:
    c.font = BOLD

wb.move_sheet('README', offset=-len(wb.sheetnames) + 1)
out = OUT / 'nabu-language-concordance.xlsx'
wb.save(out)
print(f'wrote {out}')
for k, v in counts.items():
    print(f'{k}: {v:,}')
